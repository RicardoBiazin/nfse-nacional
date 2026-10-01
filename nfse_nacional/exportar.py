"""Exportação: XMLs em pasta e planilha .xlsx das notas."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from .banco import Banco

COLUNAS = [
    ("Origem", "origem"), ("Situação", "situacao"), ("Número", "numero"),
    ("Série", "serie"), ("nDPS", "n_dps"), ("Competência", "d_compet"),
    ("Emissão", "dh_emissao"), ("Prestador CNPJ/CPF", "prestador_doc"),
    ("Prestador", "prestador_nome"), ("Tomador CNPJ/CPF", "tomador_doc"),
    ("Tomador", "tomador_nome"), ("Descrição", "descricao"),
    ("Valor serviço", "v_serv"), ("Valor líquido", "v_liq"), ("ISSQN", "v_issqn"),
    ("Chave de acesso", "chave"), ("Mensagem", "mensagem"),
]
NUMERICAS = {"v_serv", "v_liq", "v_issqn"}


def exportar_xlsx(linhas, destino: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"
    ws.append([c[0] for c in COLUNAS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for ln in linhas:
        linha = []
        for _, campo in COLUNAS:
            v = ln[campo]
            if campo in NUMERICAS and v not in (None, ""):
                v = float(Decimal(v))
            linha.append(v)
        ws.append(linha)
    for i, (titulo, campo) in enumerate(COLUNAS, 1):
        if campo in NUMERICAS:
            for cel in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in cel:
                    c.number_format = "#,##0.00"
        ws.column_dimensions[get_column_letter(i)].width = 48 if campo == "chave" else max(12, len(titulo) + 4)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    destino = Path(destino)
    wb.save(destino)
    return destino


def exportar_xmls(banco: Banco, ids: list[int], pasta: str | Path) -> int:
    """Grava o XML da NFS-e (ou a DPS, se não houver nota) e os eventos de cada nota."""
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    n = 0
    for id_ in ids:
        nota = banco.nota(id_)
        if not nota:
            continue
        if nota["nfse_xml"]:
            (pasta / f"NFSe_{nota['chave']}.xml").write_bytes(nota["nfse_xml"])
            n += 1
        elif nota["dps_xml"]:
            (pasta / f"{nota['id_dps']}.xml").write_bytes(nota["dps_xml"])
            n += 1
        if nota["chave"]:
            for ev in banco.eventos(nota["chave"]):
                (pasta / f"Evento_{nota['chave']}_{ev['tipo']}.xml").write_bytes(ev["xml"])
    return n
